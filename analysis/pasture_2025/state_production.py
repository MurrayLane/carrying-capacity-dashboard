#!/usr/bin/env python3
"""SCOPE.md E1, first test: the model gives each zone the national meat production in proportion to its dry sheep equivalents (DSE), so a state's meat per DSE is the same everywhere.
Does the ABS record of meat produced by state agree?

Usage: python3 -I analysis/pasture_2025/state_production.py ABS_DIR OUT.csv          (repository root; ABS_DIR holds 7215009 to 7215012.xlsx and the AALDC herd and flock series, as for livestock_zone.py)
For the three years 2022-23 to 2024-25: the state's share of national DSE (sheep, and beef cattle, from the 2010 class mix scaled by the state's herd or flock, as the model does) against the state's share
of the tonnes of meat produced (ABS Livestock Products, cattle and calves; sheep and lambs; states of slaughter), and the ratio, which is the state's meat per DSE relative to the nation's."""
import csv, datetime, json, os, re, sys
import openpyxl
ROOT = os.getcwd(); ABS, OUT = sys.argv[1:3]
P = json.load(open(os.path.join(ROOT, 'pipeline/yields/archive_livestock_parameters.json')))['parameters']
base = json.load(open(os.path.join(ROOT, 'pipeline/yields/archive_livestock_inputs.json')))['years']['2010']['regions']
YEARS = ['2022-23', '2023-24', '2024-25']; BASE = '2009-10'
STATES = {2: ('NSW-ACT', 1), 15: ('Victoria', 2), 26: ('Queensland', 3), 40: ('South Australia', 4), 49: ('Western Australia', 5), 56: ('Tasmania', 6), 60: ('Northern Territory', 7)}
def series(file, item_re):
    ws = openpyxl.load_workbook(os.path.join(ABS, file), read_only=True, data_only=True)['Table 1']; rows = list(ws.iter_rows(values_only=True)); cols = {}
    for c, h in enumerate(rows[3]):
        m = re.match(r'\s*(\d{4}-\d{2})', str(h))
        if m: cols[m.group(1)] = c
    return {int(r[1]): {y: r[c] for y, c in cols.items() if isinstance(r[c], (int, float))} for r in rows[4:] if r[2] and re.search(item_re, str(r[2]))}
beef, sheep = series('AALDC_Cattle_herd_series_2005_to_2025.xlsx', r'^Beef cattle'), series('AALDC_Sheep_flock_series_2005_to_2025.xlsx', r'^Sheep')
ratio = lambda s, st, y: 1.0 if not s.get(st, {}).get(BASE) or s[st].get(y) is None else s[st][y] / s[st][BASE]
def dse(z, year):
    g = lambda k: base[str(z)].get(k) or 0.0; stc = STATES[z][1]
    sh = (P['dseLambs'] * g('sheepLambs') + P['dseEwesToLamb'] * g('sheepEwesToLamb') + P['dseEwesNotInLamb'] * g('sheepEwesNotInLamb') + P['dseRams'] * g('sheepRams')) * ratio(sheep, stc, year)
    be = (P['dseBeefBulls'] * g('beefBulls') + P['dseBeefCows'] * g('beefCows') + P['dseBeefCalves'] * g('beefCalves') + P['dseBeefBalance'] * g('beefBalance')) * ratio(beef, stc, year)
    return sh, be
def monthly(file, col):
    ws = openpyxl.load_workbook(os.path.join(ABS, file), read_only=True, data_only=True)['Data1']; rows = list(ws.iter_rows(values_only=True)); out = {}
    for row in rows[10:]:
        if isinstance(row[0], datetime.datetime) and isinstance(row[col], (int, float)): out[row[0]] = row[col]
    return out
def header(file):
    ws = openpyxl.load_workbook(os.path.join(ABS, file), read_only=True, data_only=True)['Data1']; return [str(x) for x in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
STATE_NAME = {'New South Wales': 2, 'Victoria': 15, 'Queensland': 26, 'South Australia': 40, 'Western Australia': 49, 'Tasmania': 56, 'Northern Territory': 60}
def by_state(files):
    out = {z: {y: 0.0 for y in YEARS} for z in STATES}
    for f in files:
        h = header(f)
        for c, name in enumerate(h):
            for sn, z in STATE_NAME.items():
                if f'; {sn} ;' in name or name.rstrip().endswith(f'{sn} ;'):
                    d = monthly(f, c)
                    for y in YEARS:
                        fyr = int('20' + y[-2:]); out[z][y] += sum(v for dt, v in d.items() if datetime.datetime(fyr - 1, 7, 1) <= dt <= datetime.datetime(fyr, 6, 30))
    return out
cattle, sheepm = by_state(['7215009.xlsx', '7215010.xlsx']), by_state(['7215011.xlsx', '7215012.xlsx'])
rows = [['species', 'zone', 'state', 'year', 'dse_share_pct', 'meat_tonnes', 'meat_share_pct', 'meat_per_dse_relative']]
for sp, meat, k in (('beef cattle', cattle, 1), ('sheep', sheepm, 0)):
    for y in YEARS:
        tot_dse = sum(dse(z, y)[k] for z in STATES); tot_m = sum(meat[z][y] for z in STATES)
        for z, (name, _) in STATES.items():
            ds, m = dse(z, y)[k] / tot_dse, meat[z][y] / tot_m
            rows.append([sp, z, name, y, round(100 * ds, 2), round(meat[z][y]), round(100 * m, 2), round(m / ds, 3) if ds else ''])
csv.writer(open(OUT, 'w', newline='')).writerows(rows)
for r in rows[1:]:
    if r[3] == '2024-25': print(r[0].ljust(12), r[2].ljust(18), 'DSE', str(r[4]).rjust(6), '% meat', str(r[6]).rjust(6), '% ratio', r[7])
